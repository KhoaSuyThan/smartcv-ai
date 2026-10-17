from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

out=Path('scratch/quiz_plan_output/Plan_nang_cap_trac_nghiem_SmartCV.docx')
out.parent.mkdir(parents=True,exist_ok=True)
doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(.68); sec.bottom_margin=Inches(.68); sec.left_margin=Inches(.78); sec.right_margin=Inches(.78)
styles=doc.styles
normal=styles['Normal']; normal.font.name='Aptos'; normal.font.size=Pt(10); normal.font.color.rgb=RGBColor(35,48,66)
normal.paragraph_format.space_after=Pt(4); normal.paragraph_format.line_spacing=1.08
for nm,size in [('Title',25),('Heading 1',16),('Heading 2',12)]:
 s=styles[nm]; s.font.name='Aptos Display' if nm=='Title' else 'Aptos'; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor(12,28,53); s.paragraph_format.space_before=Pt(10 if nm!='Title' else 0); s.paragraph_format.space_after=Pt(5)

def shade(cell,fill):
 tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)
def set_cell(cell,text,bold=False,color=None,size=9):
 cell.text=''; p=cell.paragraphs[0]; p.paragraph_format.space_after=Pt(0); r=p.add_run(text); r.bold=bold; r.font.name='Aptos'; r.font.size=Pt(size)
 if color: r.font.color.rgb=RGBColor(*color)
 cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def bullet(text,level=0):
 p=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2'); p.paragraph_format.space_after=Pt(2); p.add_run(text); return p
def num(text):
 p=doc.add_paragraph(style='List Number'); p.paragraph_format.space_after=Pt(2); p.add_run(text); return p

p=doc.add_paragraph(style='Title'); p.add_run('Kế hoạch nâng cấp phần trắc nghiệm SmartCV AI')
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(9); r=p.add_run('Đề xuất kỹ thuật dựa trên mã nguồn hiện tại và giao diện Practice'); r.italic=True; r.font.color.rgb=RGBColor(84,101,122); r.font.size=Pt(11)

# Executive summary
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(8)
r=p.add_run('Kết luận. '); r.bold=True
p.add_run('Plan Gemini đúng hướng về Local AI, câu hỏi theo JD, lời giải và trải nghiệm thi. Tuy nhiên cần sửa thứ tự và một số giả định trước khi triển khai: code quiz hiện gọi Gemini trực tiếp, trang làm bài đang nhận đáp án đúng trong model, dữ liệu hiện hữu đủ chứa giải thích trong Feedback, và phương án Gemini fallback không thể được gọi là chạy offline.')

# status table
doc.add_heading('1. Hiện trạng đã xác minh',level=1)
t=doc.add_table(rows=1,cols=3); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Light Shading Accent 1'
for c,txt in zip(t.rows[0].cells,['Khu vực','Hiện trạng trong hệ thống','Hệ quả cho plan']): set_cell(c,txt,True,(255,255,255),9); shade(c,'16345B')
rows=[
('Phiên và dữ liệu','InterviewType=1; InterviewSession liên kết ResumeID và JobID. Câu hỏi/đáp án/lựa chọn nằm trong InterviewMessages.','Tái sử dụng cấu trúc hiện tại.'),
('Sinh đề','StartInterview xây CV summary nhưng nhánh quiz gọi _aiService.GenerateContent (Gemini), chỉ gửi CV.','Chuyển sang Ollama trước; thêm nguồn JD rõ ràng.'),
('Đáp án','SelectedAnswer của bản ghi interviewer đang chứa đáp án đúng. PracticeQuiz bind cả session.Messages.','Ẩn đáp án khỏi model GET; chỉ dùng ở service chấm phía server.'),
('Nộp bài','SubmitQuizAnswer lưu từng lựa chọn vào bản ghi candidate và Feedback chỉ ghi đúng/sai.','Bổ sung explanation và chủ đề sai vào phản hồi sau khi nộp.'),
('Giao diện','Practice có lựa chọn nguồn/loại/số câu; PracticeQuiz hiển thị toàn bộ câu trên một trang, chưa có timer/palette.','Tái sử dụng lựa chọn hiện có; bổ sung điều hướng và thời gian.'),
('JD','JobID được liên kết/tự suy luận cho phiên; prompt quiz chưa đưa nội dung JD.','Đọc Job theo JobID đã chọn, kiểm tra quyền và chỉ lấy trường cần thiết.'),
]
for row in rows:
 cells=t.add_row().cells
 for c,txt in zip(cells,row): set_cell(c,txt,size=8.5)

doc.add_heading('2. Đánh giá plan Gemini',level=1)
for s in [
'Đúng: dùng Local Ollama JSON Mode làm hướng chính; có template dự phòng; yêu cầu giải thích sau khi trả lời; kiểm thử cả nhánh CV và JD; thiết kế timer, question palette, Dark Mode và tư vấn điểm yếu.',
'Cần sửa: ghi “Ollama → Gemini → template” là failover chung không đồng nghĩa offline. Nếu cần chạy offline tuyệt đối, dùng Ollama → template. Gemini chỉ là fallback tùy chọn khi cấu hình cho phép, mạng hoạt động và dữ liệu đã được giảm thiểu/ẩn danh theo policy.',
'Cần sửa: JSON mẫu phải được validate nghiêm ngặt trước khi lưu: đủ đúng số câu, 4 lựa chọn duy nhất, correctAnswer thuộc A–D, không trùng lựa chọn, explanation không rỗng; thử sinh lại giới hạn một lần rồi dùng template.',
'Cần sửa: unverifiedSkills/missingSkills không có sẵn đồng nhất trong mọi luồng. Tạo tập trọng tâm từ JD trước, đối chiếu CV như tín hiệu ưu tiên, không kết luận kỹ năng “thiếu” chỉ vì parser không tìm thấy.',
'Cần sửa: dự toán 1–1,5 ngày là ước tính lạc quan cho 4 phase gồm backend, bảo mật đáp án, JD, UI, lưu kết quả, timer phía server và test. Nên lập theo mốc đầu ra, thời lượng xác nhận sau spike Local AI/JD và kiểm tra schema.',
]: bullet(s)

doc.add_heading('3. Mục tiêu và nguyên tắc triển khai',level=1)
for s in [
'Sinh bộ trắc nghiệm chuyên môn dựa trên CV hoặc JD cụ thể bằng Ollama cục bộ, lưu lại trong InterviewMessages và tiếp tục được phiên nếu tải lại trang.',
'Chấm điểm hoàn toàn phía server. Không gửi correctAnswer hoặc explanation tới trình duyệt trong phiên đang làm.',
'Giải thích câu đúng/sai chỉ hiện sau khi nộp; lời khuyên tổng hợp dựa trên topic/skill gắn với câu sai, tránh yêu cầu LLM chấm lại đáp án khách quan.',
'Timer phải được xác thực server-side bằng thời điểm bắt đầu và thời lượng đề. JavaScript chỉ hiển thị đếm ngược, không quyết định phiên còn hạn.',
'Không đổi schema trong phạm vi cơ sở. Tái dùng Feedback của câu hỏi để lưu JSON có version gồm explanation/topic, hoặc dùng AiEvaluation cho metadata phiên nếu trường hiện tại phù hợp. Chỉ đề xuất migration riêng nếu spike xác nhận cách này không an toàn/khả thi; khi đó cần duyệt trước và đồng bộ DtbDoAnCS cùng Program.cs.',
]: bullet(s)

doc.add_heading('4. Kiến trúc sinh đề và hợp đồng dữ liệu',level=1)
doc.add_heading('4.1. Quy trình nguồn nội dung',level=2)
for s in [
'Theo CV: Resume.JsonContent → trích kỹ năng/kinh nghiệm liên quan → loại thông tin định danh khỏi prompt → sinh câu hỏi kiểm tra kiến thức của kỹ năng đó, không hỏi lại thông tin tiểu sử.',
'Theo JD: JobID đã chọn → kiểm tra tồn tại/trạng thái và quyền truy cập → lấy tiêu đề, mô tả công việc, yêu cầu và kỹ năng → giới hạn độ dài → ưu tiên kỹ năng bắt buộc. CV chỉ dùng để cá nhân hóa độ khó nếu đã chọn.',
'Nếu JD có ít dữ kiện, hỏi các kiến thức nền phổ biến của nhóm kỹ năng; không bịa yêu cầu riêng của doanh nghiệp.',
]: bullet(s)
doc.add_heading('4.2. Hợp đồng JSON chuẩn hóa',level=2)
code='''{\n  "question": "Câu hỏi có một đáp án tốt nhất",\n  "options": ["A. ...", "B. ...", "C. ...", "D. ..."],\n  "correctAnswer": "B",\n  "explanation": "Vì sao đáp án đúng và vì sao các lựa chọn dễ nhầm không phù hợp",\n  "topic": "React Hooks"\n}'''
p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.2); p.paragraph_format.space_after=Pt(5); r=p.add_run(code); r.font.name='Consolas'; r.font.size=Pt(9); r.font.color.rgb=RGBColor(28,55,85)
for s in [
'Chuẩn hóa tên field không phân biệt hoa thường; từ chối câu hỏi thiếu field hoặc loại dữ liệu sai. Kiểm tra đúng số câu và chuẩn đáp án A–D trước khi ghi DB.',
'Lưu question vào Content, options vào ChoicesJson và correctAnswer vào SelectedAnswer của câu hỏi. Lưu explanation/topic trong Feedback của bản ghi câu hỏi dưới dạng JSON có version. Trong GET phiên đang làm, map sang QuizViewModel loại bỏ SelectedAnswer và Feedback.',
'Không ghi CV/JD đầy đủ hoặc thông tin cá nhân vào log prompt. Log provider, model, latency, số token nếu có, lỗi parse và mã phiên; không log API key.',
]: bullet(s)

doc.add_heading('5. Các giai đoạn triển khai',level=1)
phases=[
('Giai đoạn 0 · Chốt hợp đồng và spike kỹ thuật','Đọc trọn luồng StartInterview, binding Practice source/JobID, quyền sở hữu Job/Resume và nội dung Job/Resume.JsonContent. Xác nhận cách lưu explanation trong Feedback và thời lượng đề trong metadata hiện có. Thử JSON Mode Ollama với Qwen 2.5 3B/1.5B trên 5 câu.','Quyết định nguồn dữ liệu, DTO, mẫu template, budget token/timeout và ranh giới cloud. Không sửa schema.'),
('Giai đoạn 1 · Backend Local AI và failover','Tách quiz generator khỏi action/controller thành service có provider result. Ollama là mặc định. Parse/validate JSON, yêu cầu sửa/sinh lại tối đa một lần. Khi Ollama lỗi, dùng bộ template kỹ thuật offline theo skill. Gemini là fallback tùy chọn sau template hoặc trước template chỉ khi bật cấu hình, có mạng, policy cho phép và payload đã được giảm thiểu. Dùng transaction/lưu session chỉ khi sinh đủ số câu hợp lệ; xóa phiên lỗi an toàn.','Sinh đề CV/JD chạy offline; mỗi câu lưu options, đáp án đúng, explanation, topic; không để phiên rỗng/nửa chừng.'),
('Giai đoạn 2 · Theo JD và chẩn đoán kỹ năng','Thêm/hoàn thiện lựa chọn source CV/JD trong Practice; gắn JobID/ResumeID vào request; xác thực ownership và giới hạn input. Ưu tiên kỹ năng trong JD; so CV với JD để xếp thứ tự, ghi rõ “không thấy trong CV” thay vì kết luận thiếu năng lực. Phân bố câu theo kỹ năng và mức độ.','Cả CV-only và JD-targeted sinh đúng nội dung; metadata đề ghi nguồn và kỹ năng trọng tâm để hiển thị lịch sử.'),
('Giai đoạn 3 · Màn hình thi','Tạo QuizViewModel an toàn; mỗi thời điểm hiện một câu hoặc nhóm câu gọn. Bổ sung palette đã trả lời/chưa trả lời/đang xem, quay lại câu, trạng thái tương phản đủ cả light/dark, lưu tiến độ lựa chọn nếu refresh bằng sessionStorage hoặc POST nháp nếu cần. Countdown hiển thị theo server-issued start/deadline; server tự khóa và chấm khi hết giờ. Hộp xác nhận trước khi nộp thủ công nếu còn câu trống.','Có thể điều hướng, biết câu bỏ trống; refresh không reset thời gian; hết giờ không thể kéo dài bằng sửa JavaScript.'),
('Giai đoạn 4 · Kết quả và tư vấn','Chấm một lần theo correctAnswer phía server; idempotent khi bấm kép hoặc gửi lại. Lưu candidate answer theo MessageID của câu thay vì phụ thuộc thứ tự thời gian. Sau nộp mới hiện đúng/sai, đáp án chuẩn, explanation và topic. Thống kê số câu đã trả lời/bỏ trống, điểm và thời gian; tư vấn tổng hợp có thể dùng template cục bộ trước, LLM chỉ để diễn đạt nếu latency cho phép.','Trang kết quả có giải thích theo câu, số liệu phiên ổn định và gợi ý ôn tập theo nhóm lỗi.'),
('Giai đoạn 5 · Kiểm thử và hoàn thiện','Kiểm tra unit cho parser/validator/scorer; luồng Ollama sẵn, Ollama thiếu model, timeout, JSON sai, template fallback, Gemini bị tắt; CV/JD hợp lệ và không có quyền; câu thiếu đáp án; Free/Pro min/max; nộp kép, hết giờ, refresh, Dark Mode và lịch sử. Chạy build cùng test/API phù hợp sau khi code được triển khai.','Biên bản kiểm thử, lỗi đã xử lý, latency/độ ổn định và hướng dẫn bật Ollama/model.'),
]
for title,desc,deliver in phases:
 doc.add_heading(title,level=2)
 p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(2); r=p.add_run('Công việc. '); r.bold=True; p.add_run(desc)
 p=doc.add_paragraph(); p.paragraph_format.left_indent=Inches(.18); p.paragraph_format.space_after=Pt(4); r=p.add_run('Đầu ra. '); r.bold=True; r.font.color.rgb=RGBColor(31,112,91); p.add_run(deliver)

doc.add_heading('6. Quy tắc thời gian, giới hạn và trạng thái phiên',level=1)
for s in [
'Giữ giới hạn đang có ở backend: Free 5–10 câu, Pro 5–20 câu. Giao diện chỉ hỗ trợ nhập liệu; backend tiếp tục validate và trả thông báo rõ ràng.',
'Đề xuất thời lượng 90 giây/câu, cấu hình được và chốt sau spike; 5 câu = 7 phút 30 giây, 20 câu = 30 phút. Đây là cấu hình đề xuất, không phải ràng buộc nghiệp vụ đã xác nhận.',
'Không tính thời gian sinh đề vào thời gian làm bài. Deadline bắt đầu khi bộ câu hỏi được lưu thành công. Khi nộp, server so sánh thời gian hiện tại với deadline và coi câu không chọn là bỏ trống.',
'Nếu chưa có CompletedAt cho phiên đang làm thì khi hết giờ lưu thời điểm kết thúc; nếu bắt buộc phải lưu deadline bền vững qua server restart mà metadata hiện có không đủ, đây là điểm cần trình người dùng duyệt schema trước khi làm migration.',
]: bullet(s)

doc.add_heading('7. Tiêu chí nghiệm thu',level=1)
checks=[
('Sinh đề','Số câu thực tế đúng yêu cầu; mọi câu có 4 lựa chọn A–D, đáp án hợp lệ, giải thích và topic.'),
('Offline','Tắt mạng/Gemini nhưng Ollama hoạt động: sinh và chấm đề vẫn hoàn tất qua template nếu LLM lỗi.'),
('Bảo mật đáp án','Trước submit, HTML/JSON response của trang không chứa correctAnswer hoặc explanation; chấm chỉ diễn ra server-side.'),
('Theo JD','Đề ưu tiên các kỹ năng được trích từ đúng JobID đã chọn; không dùng JD không thuộc quyền truy cập.'),
('Timer','Refresh không làm mới thời gian; quá deadline thì server khóa nộp thường và chấm câu đã chọn, câu còn lại để trống.'),
('Kết quả','Mỗi câu hiện lựa chọn của người dùng, đáp án chuẩn, lời giải; tổng điểm/tỷ lệ trả lời khớp số câu.'),
('Tương thích','Lịch sử InterviewType=1 hiện có vẫn mở được; model dữ liệu cũ thiếu explanation được hiển thị thông báo phù hợp, không crash.'),
('UI','Palette điều hướng đúng; trạng thái focus/đã chọn/chưa chọn nhìn rõ ở light/dark và màn hình nhỏ.'),
]
t=doc.add_table(rows=1,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Light Shading Accent 1'
for c,txt in zip(t.rows[0].cells,['Tiêu chí','Điều kiện đạt']): set_cell(c,txt,True,(255,255,255),9); shade(c,'16345B')
for a,b in checks:
 cells=t.add_row().cells; set_cell(cells[0],a,True,size=8.5); set_cell(cells[1],b,size=8.5)

doc.add_heading('8. Rủi ro và cách xử lý',level=1)
risks=[
('Ollama chưa cài hoặc chưa tải model','Hướng dẫn cấu hình rõ ràng; template offline là fallback chức năng; không xóa session trước khi có phương án thay thế.'),
('LLM bịa đáp án/JSON sai','Schema validation, kiểm tra đáp án, template đã review; đo lỗi và tỷ lệ fallback trên bộ đề mẫu.'),
('Rò rỉ đáp án','ViewModel tách biệt; không bind entity trực tiếp; kiểm tra HTML/response network trước submit.'),
('Cloud fallback lộ CV/JD','Tắt mặc định hoặc yêu cầu cấu hình opt-in; bỏ PII, giảm payload; nếu policy không cho phép thì bỏ qua Gemini và dùng template.'),
('Timer lệch hoặc bị chỉnh phía client','Server deadline là nguồn quyết định; JavaScript chỉ trình bày đồng hồ.'),
('JD ngắn/không có kỹ năng','Báo rõ nguồn thiếu dữ liệu; chuyển sang kiến thức nền phù hợp, không khẳng định câu hỏi sát JD.'),
]
t=doc.add_table(rows=1,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Light Shading Accent 1'
for c,txt in zip(t.rows[0].cells,['Rủi ro','Xử lý']): set_cell(c,txt,True,(255,255,255),9); shade(c,'16345B')
for a,b in risks:
 cells=t.add_row().cells; set_cell(cells[0],a,True,size=8.5); set_cell(cells[1],b,size=8.5)

doc.add_heading('9. Ước lượng và thứ tự ưu tiên',level=1)
p=doc.add_paragraph('Ước lượng ban đầu: 3–5 phiên tập trung, khoảng 10–16 giờ thực thi và kiểm tra sau spike. Đây là khoảng dự toán thay cho cam kết “1–1,5 ngày”; thời gian thực tế phụ thuộc chất lượng output Ollama, định dạng trường Job và quyết định có cần lưu deadline bền vững hay không.')
for s in [
'P0 bắt buộc: bảo vệ đáp án, Ollama + template offline, validation, CV/JD, lưu giải thích, chấm phía server, kết quả, giới hạn Free/Pro, test luồng.',
'P1 trải nghiệm: palette, countdown, Dark Mode, giữ lựa chọn khi refresh, báo câu bỏ trống.',
'P2 nâng cao: phân phối độ khó/thời lượng theo topic, báo cáo điểm yếu nâng cao, benchmark so sánh model và Gemini opt-in.',
]: bullet(s)

doc.add_heading('10. Quyết định cần chốt trước khi code',level=1)
for s in [
'Timer dùng 90 giây/câu hay thời lượng cố định theo toàn đề? Mốc 90 giây/câu đang là giả định đề xuất.',
'Gemini fallback có được bật không, và quy định nào cho phép gửi dữ liệu đã ẩn danh lên dịch vụ ngoài? Mặc định đề xuất là tắt.',
'Có bắt buộc giữ timer chính xác qua restart backend không? Nếu có và metadata hiện tại không hỗ trợ, cần plan migration riêng để duyệt trước.',
'Nguồn JD cần dùng trường nào làm nội dung chuẩn: mô tả, yêu cầu, kỹ năng hay tổ hợp? Chốt theo schema Job sau khi kiểm tra model.',
'Có cần lưu tiến độ từng câu server-side khi đang làm hay sessionStorage là đủ? Điều này ảnh hưởng khả năng tiếp tục trên thiết bị khác.',
]: bullet(s)

# footer
for section in doc.sections:
 f=section.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.RIGHT; r=f.add_run('SmartCV AI  |  Plan nâng cấp Trắc nghiệm  |  17/10/2026'); r.font.name='Aptos'; r.font.size=Pt(8); r.font.color.rgb=RGBColor(110,124,141)

doc.core_properties.title='Kế hoạch nâng cấp phần trắc nghiệm SmartCV AI'
doc.core_properties.subject='Đánh giá và kế hoạch triển khai dựa trên mã nguồn hiện tại'
doc.core_properties.author='Nhóm SmartCV AI'
doc.save(out)
print(out)
